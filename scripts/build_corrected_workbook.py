#!/usr/bin/env python
"""Produce a corrected copy of the supplementary workbook.

The delivered workbook is not modified. This writes a separate corrected copy,
preserving its formatting, sheet names, column order and every unrelated author
edit: sheets are edited in place rather than rebuilt, so number formats, column
widths and the Contents notes survive.

    python scripts/build_corrected_workbook.py

Source of truth: the regenerated CSVs under outputs/tables/publication/. The
script refuses to run if those are absent, and verifies cell by cell afterwards
that the corrected workbook and the CSVs agree.
"""
from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "publication_artwork" / "tables" / "SLA_Supplementary_Tables.xlsx"
CSVS = ROOT / "outputs" / "tables" / "publication"
DEST = ROOT / "outputs" / "tables" / "SLA_Supplementary_Tables_corrected.xlsx"

SHEET_CSV = {
    "S1": "TableS1_haplotype_nearest_and_most_distant_HLA.csv",
    "S2": "TableS2_canonical_split_leakage_and_five_seed_validation.csv",
    "S3": "TableS3_HLA_eplet_associated_position_statistics.csv",
    "S4a": "TableS4a_structural_model_confidence_per_model.csv",
    "S4b": "TableS4b_structural_superpositions_per_pair.csv",
    "S5": "TableS5_structural_position_annotations.csv",
    "S6": "TableS6_experimental_validation_priorities.csv",
    "Dictionary": "PUBLICATION_TABLES_DATA_DICTIONARY.csv",
}
NEW_SHEET = ("S4b_provenance", "TableS4b_provenance_submitted_version.csv")

#: Columns renamed between the approved workbook and the corrected one. Each
#: inherits the style and number format of the column it replaces, by name.
RENAMED = {
    "hla_allele_in_submitted_legend": "hla_allele_in_legend",
    "allele_match_vs_submitted_legend": "allele_match",
    "rmsd_printed_in_submitted_figure": "rmsd_reported_in_manuscript",
    "delta_cealign_minus_submitted": "delta_cealign_minus_reported",
    "delta_super_minus_submitted": "delta_super_minus_reported",
    "delta_align_minus_submitted": "delta_align_minus_reported",
}

#: Explicit formats for columns this repository adds, which have no counterpart
#: in the approved workbook to inherit from.
NEW_FORMATS = {
    "record_level": "General",
    "identity_method": "General",
    "displayed_equals_round1_of_cealign": "General",
    "note": "General",
    "haplotypes_all_curated": "General",
    "haplotype_selection_source": "General",
    # what the panel prints: one decimal, as the figure shows it
    "rmsd_displayed_in_final_figure": "0.0",
}

#: Fallbacks when a column is neither known nor renamed. Counts are integers;
#: everything measured keeps four decimals, which is what the approved workbook
#: uses throughout.
FALLBACK_INT = "0"
FALLBACK_FLOAT = "0.0000"


def _column_styles(ws) -> dict:
    """{header: (header style, body style, number format)} from a sheet, by name.

    Keyed by name rather than by position. The previous version captured styles
    positionally, so when Table S4b lost six columns and gained four, every style
    after the first change landed on the wrong column: an integer format fell on
    `percent_identity_over_shared_positions` and displayed 75.58 as 76, and on
    `coverage_of_sla_chain_ATOMS`, where 0.646 displayed as 1. It also keyed the
    format map on the second row's *values* instead of the header.
    """
    out = {}
    head = list(ws[1])
    body = list(ws[2]) if ws.max_row > 1 else head
    for i, h in enumerate(head):
        name = str(h.value) if h.value is not None else None
        if name is None:
            continue
        b = body[i] if i < len(body) else h
        out[name] = (copy(h._style), copy(b._style), b.number_format)
    return out


def _resolve(name: str, styles: dict, series: pd.Series, default):
    """Style and number format for one output column, chosen by name."""
    if name in styles:
        return styles[name]
    if name in RENAMED and RENAMED[name] in styles:
        head, body, _fmt = styles[RENAMED[name]]
        return head, body, NEW_FORMATS.get(name, _fmt)
    head, body, _ = default
    if name in NEW_FORMATS:
        return head, body, NEW_FORMATS[name]
    if pd.api.types.is_bool_dtype(series):
        return head, body, "General"
    if pd.api.types.is_integer_dtype(series):
        return head, body, FALLBACK_INT
    if pd.api.types.is_float_dtype(series):
        return head, body, FALLBACK_FLOAT
    return head, body, "General"


def _write_sheet(ws, df: pd.DataFrame, styles: dict | None = None) -> None:
    """Replace a sheet's contents, carrying styles and formats across BY NAME."""
    styles = _column_styles(ws) if styles is None else styles
    default = (copy(ws[1][0]._style),
               copy((ws[2] if ws.max_row > 1 else ws[1])[0]._style), "General")
    ws.delete_rows(1, ws.max_row)
    resolved = [_resolve(col, styles, df[col], default) for col in df.columns]
    for j, (col, (head, _b, _f)) in enumerate(zip(df.columns, resolved), start=1):
        ws.cell(row=1, column=j, value=col)._style = head
    for i, (_, row) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            ws.cell(row=i, column=j,
                    value=None if pd.isna(v) else (v.item() if hasattr(v, "item") else v))
    # Styles and formats go on last: writing a value can reset a cell's format,
    # so applying them first silently loses them for the added columns.
    for j, (head, body, fmt) in enumerate(resolved, start=1):
        for i in range(2, len(df) + 2):
            cell = ws.cell(row=i, column=j)
            cell._style = copy(body)
            cell.number_format = fmt


def main() -> int:
    if not SRC.exists():
        raise SystemExit(f"authoritative workbook not found: {SRC}")
    missing = [n for n in SHEET_CSV.values() if not (CSVS / n).exists()]
    if missing:
        raise SystemExit("run `make reproduce` first; missing regenerated CSVs: "
                         + ", ".join(missing))
    DEST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, DEST)
    wb = openpyxl.load_workbook(DEST)

    s4b_styles = _column_styles(wb["S4b"]) if "S4b" in wb.sheetnames else {}
    for sheet, csv in SHEET_CSV.items():
        if sheet not in wb.sheetnames:
            print(f"  sheet {sheet} absent from the workbook; skipped")
            continue
        _write_sheet(wb[sheet], pd.read_csv(CSVS / csv))
        print(f"  {sheet}: updated from {csv}")

    name, csv = NEW_SHEET
    if (CSVS / csv).exists():
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name, wb.sheetnames.index("S4b") + 1)
        # A new sheet has no styles of its own: it inherits S4b's, by name, so
        # the renamed submitted-version columns keep their original formats.
        ws.cell(row=1, column=1, value="seed")
        ws.cell(row=2, column=1, value=None)
        if s4b_styles:
            head, body, _ = next(iter(s4b_styles.values()))
            ws.cell(row=1, column=1)._style = head
            ws.cell(row=2, column=1)._style = body
        _write_sheet(ws, pd.read_csv(CSVS / csv), styles=s4b_styles)
        print(f"  {name}: added from {csv}")

    # Contents: refresh row counts and register the new sheet
    contents = wb["Contents"]
    rows = {contents.cell(row=r, column=1).value: r for r in range(1, contents.max_row + 1)}
    for sheet, csv in SHEET_CSV.items():
        if sheet in rows:
            contents.cell(row=rows[sheet], column=2,
                          value=int(len(pd.read_csv(CSVS / csv))))
    if "S4b" in rows and NEW_SHEET[0] not in rows:
        at = rows["S4b"] + 1
        contents.insert_rows(at)
        for j, v in enumerate([NEW_SHEET[0], int(len(pd.read_csv(CSVS / NEW_SHEET[1]))),
                               "Submitted-version record for S4b (historical; not current values)",
                               NEW_SHEET[1]], start=1):
            c = contents.cell(row=at, column=j, value=v)
            c._style = copy(contents.cell(row=rows["S4b"], column=j)._style)
        print("  Contents: S4b_provenance registered")

    wb.save(DEST)

    # --- verify the corrected workbook against the CSVs, cell by cell
    wb2 = openpyxl.load_workbook(DEST, data_only=True)
    problems = []
    for sheet, csv in list(SHEET_CSV.items()) + [NEW_SHEET]:
        if sheet not in wb2.sheetnames:
            continue
        df = pd.read_csv(CSVS / csv)
        ws = wb2[sheet]
        got = pd.DataFrame(ws.iter_rows(min_row=2, values_only=True),
                           columns=[c.value for c in ws[1]])
        if list(got.columns) != list(df.columns):
            problems.append(f"{sheet}: column mismatch")
            continue
        if len(got) != len(df):
            problems.append(f"{sheet}: {len(got)} rows vs {len(df)}")
            continue
        for c in df.columns:
            if pd.api.types.is_numeric_dtype(df[c]):
                a = pd.to_numeric(got[c], errors="coerce").astype(float)
                b = df[c].astype(float)
                if not ((a - b).abs().fillna(0) <= 5e-9).all():
                    problems.append(f"{sheet}.{c}: numeric mismatch")
            else:
                a = got[c].where(got[c].notna(), "").astype(str)
                b = df[c].where(df[c].notna(), "").astype(str)
                if (a.values != b.values).any():
                    problems.append(f"{sheet}.{c}: value mismatch")
    print(f"\nworkbook <-> CSV check: {'AGREE' if not problems else 'PROBLEMS'}")
    for p in problems:
        print("   ", p)
    print(f"\nwrote {DEST}")
    print(f"original left untouched: {SRC}")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
